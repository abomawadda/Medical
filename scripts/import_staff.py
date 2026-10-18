"""Import the owner's provider spreadsheet as data, never as executable instructions."""
import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from staff_catalog import ROOT, IMPORT, normalize_contacts

DATE = '2026-10-18'
CITY = {'السينبلاوين': 'السنبلاوين', 'المنزله': 'المنزلة', 'اجا': 'أجا', 'مركز منية النصر': 'منية النصر'}
WORDS = {'مستشفي': 'مستشفى', 'المنصوره': 'المنصورة', 'الحكمه': 'الحكمة', 'الاعصاب': 'الأعصاب', 'المتطوره': 'المتطورة', 'لامراض': 'لأمراض', 'الكلوى': 'الكلوي', 'القسطره': 'القسطرة', 'الاوعيه': 'الأوعية', 'الاوعية': 'الأوعية', 'الباطنه': 'الباطنة', 'التخصصى': 'التخصصي', 'الاورام': 'الأورام', 'الجزيره': 'الجزيرة', 'الدولى': 'الدولي', 'الاكاديمية': 'الأكاديمية', 'الالفا': 'الألفا', 'الفا': 'ألفا', 'الاجهزه': 'الأجهزة', 'للاجهزه': 'للأجهزة', 'السمعيه': 'السمعية', 'العزبى': 'العزبي', 'موافى': 'موافي', 'الطاهره': 'الطاهرة', 'للاشعه': 'للأشعة', 'للاشعة': 'للأشعة', 'اشعة': 'أشعة', 'اسنان': 'أسنان', 'الطبيعى': 'الطبيعي', 'نبضد': 'نبض', 'مصطفي': 'مصطفى', 'اورانج': 'أورانج', 'لادارة': 'لإدارة', 'الحرية': 'الحرية', 'الحريه': 'الحرية', 'امام': 'أمام', 'اعلى': 'أعلى', 'الاسلامى': 'الإسلامي', 'بجور': 'بجوار', 'التعليمة': 'التعليمية', 'ابو': 'أبو', 'احمد': 'أحمد', 'الاول': 'الأول', 'الاهلي': 'الأهلي', 'قناه': 'قناة', 'الدور3': 'الدور 3', 'الاداريه': 'الإدارية'}
KINDS = {'مستشفى': 'hospitals', 'صيدلية': 'pharmacies', 'معمل تحاليل': 'labs', 'مركز أشعة': 'radiology', 'سمعيات': 'specialized_centers', 'علاج طبيعي': 'specialized_centers', 'أسنان': 'specialized_centers'}
ALIASES = {'758': 'hosp_sellab', '766': 'hosp_kheir', '767': 'hosp_hekma'}
HOLD = {'763': 'تعارض جغرافي: دكرنس مقابل عنوان داندي مول بطريق القاهرة الإسكندرية', '2470': 'تعارض بين اسم العبور ومدينة منية النصر وكود أرضي 055؛ يلزم تأكيد المقر'}
OUTSIDE = {'1499': 'بورسعيد', '2907': 'دمياط', '2915': 'دمياط'}
EVIDENCE = {
    '19989': ('https://technoscan.com.eg/AR/main_pages/branch_locator.aspx', 'official_publication'),
    '16064': ('https://royal-lab.net/contact-us/', 'official_publication'),
    '19600': ('https://elezabypharmacy.com/', 'official_publication'),
    '19911': ('https://www.alborglab.com/', 'official_publication'),
    '0502332691': ('https://www.140online.com/company.aspx?CompanyId=C5449', 'directory_crosscheck'),
    '0502332692': ('https://www.140online.com/company.aspx?CompanyId=C5449', 'directory_crosscheck'),
    '0502332693': ('https://www.140online.com/company.aspx?CompanyId=C5449', 'directory_crosscheck'),
}
LOGOS = {
    'تكنوسكان': ('technoscan.png', 'https://technoscan.com.eg/AR/Home_page/'),
    'البرج': ('al-borg.png', 'https://www.alborglab.com/'),
    'العزبي': ('el-ezaby.png', 'https://elezabypharmacy.com/'),
}


def clean(value):
    t = re.sub(r'\s+', ' ', str(value or '')).strip()
    t = re.sub(r'(?<!\w)ش[.\s]+', 'شارع ', t)
    t = re.sub(r'(?<=\d)ش(?=[\u0600-\u06ff])', ' شارع ', t)
    for before, after in WORDS.items():
        t = re.sub(r'(?<!\w)' + re.escape(before) + r'(?!\w)', after, t)
    return t


def sheet(book, name, headers, rows):
    s = book.create_sheet(name)
    s.sheet_view.rightToLeft = True
    s.freeze_panes = 'A2'
    s.append(headers)
    for row in rows:
        s.append(row)
    for cell in s[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=12)
        cell.fill = PatternFill('solid', fgColor='552785')
    for row in s.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Formula-like cell content remains inert text.
                cell.number_format = '@'
    s.auto_filter.ref = s.dimensions
    for column in s.columns:
        s.column_dimensions[column[0].column_letter].width = min(62, max(16, max(len(str(c.value or '')) for c in column[:120]) * .65))
    return s


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', required=True)
    parser.add_argument('--output-dir', required=True)
    args = parser.parse_args()
    source = Path(args.source)
    book = openpyxl.load_workbook(source, data_only=False)
    if book.active.max_column != 11:
        raise ValueError('Unexpected provider workbook structure; review columns before importing.')
    rows = list(book.active.values)[1:]
    rows = [r for r in rows if r[0] is not None]
    codes = [str(r[0]) for r in rows]
    if len(set(codes)) != len(codes):
        raise ValueError('Duplicate source provider codes require review.')
    contacts = normalize_contacts([str(r[10] or '') for r in rows])
    accepted, normalized, reviews, phones, corrections, original = [], [], [], [], [], []
    stats = Counter()
    for row_number, (row, contact) in enumerate(zip(rows, contacts), 2):
        code, raw_name, raw_type, raw_governorate, raw_city, raw_address, location, english, invoice, channel, raw_phone = row
        code = str(code)
        name, category, address = clean(raw_name), clean(raw_type), clean(raw_address)
        city = CITY.get(clean(raw_city), clean(raw_city))
        markaz = 'أجا' if city == 'صهرجت الصغرى' else city
        governorate = OUTSIDE.get(code, 'الدقهلية')
        status = 'خارج النطاق الحالي' if code in OUTSIDE else 'معلق للمراجعة' if code in HOLD else 'جاهز للدمج'
        issues = []
        if code in OUTSIDE:
            issues.append('المحافظة الأصلية غير صحيحة؛ محفوظ للتوسع المستقبلي ولا ينشر بالدقهلية')
        if code in HOLD:
            issues.append(HOLD[code])
        if code == '748':
            issues.append('الاسم الإنجليزي يخص مستشفى لا الصيدلية؛ استبعد من العرض')
        if code == '1928':
            issues.append('وصف البنك غير موثوق؛ حذف من العنوان المعروض فقط وحفظ في الأصل')
            address = 'ميدان المحافظة — تفاصيل المقر تحتاج تأكيدًا'
        if code == '768':
            issues.append('العنوان B7 غير كافٍ؛ راجع الخريطة وأكد العنوان قبل الزيارة')
        if code == '2745':
            issues.append('الأرضي 093 تابع لمنطقة خارج الدقهلية؛ يلزم تأكيد أنه رقم هذه الجهة')
        for invalid in contact['invalid']:
            issues.append('رقم معلق: ' + invalid['raw'] + ' — ' + invalid['reason'])
            stats['invalid_phones'] += 1
        for field, before, after in [('الاسم', raw_name, name), ('النوع', raw_type, category), ('المدينة', raw_city, city), ('العنوان', raw_address, address), ('المحافظة', raw_governorate, governorate)]:
            if str(before or '').strip() != after:
                corrections.append([code, row_number, field, str(before or ''), after, 'تصحيح موثق بالمراجعة؛ الأصل محفوظ'])
        if not isinstance(location, str) or not location.startswith(('https://maps.', 'https://goo.gl/maps/')):
            raise ValueError(f'Unexpected map URL in source row {row_number}')
        kind = 'records' if code == '773' else KINDS[category]
        entity_id = ALIASES.get(code, 'staff_' + code)
        evidence = [{ 'number': n, 'source_url': EVIDENCE[n][0], 'check_type': EVIDENCE[n][1], 'checked_at': DATE} for k in ['mobile', 'landline', 'shortcode'] for n in contact[k] if n in EVIDENCE]
        valid_phone = ' - '.join(n for k in ['mobile', 'landline', 'shortcode'] for n in contact[k])
        entity = dict(id=entity_id, name=name, address=address, city=city, markaz=markaz, country_code='EG', governorate=governorate, governorate_code='EG-DK', maps_url=location.strip(), phone=valid_phone,
                      contact_groups={k: contact[k] for k in ['mobile', 'landline', 'shortcode']}, contact_evidence=evidence, contact_review_required=bool(contact['invalid']) or code == '2745',
                      source_ref='staff-20261018', source_row=row_number, staff_provider_code=code, imported_at=DATE,
                      network_members=['employee'], contract_type='direct_we', contract_entity='تعاقد مباشر من الشركة WE', contract_entity_id='direct_we',
                      coverage_scope='unknown', declared_coverage_scope='employee_only', coverage_verified=False, coverage_label='مدرج بقائمة الموظفين؛ التغطية والخدمة تحتاجان تأكيدًا',
                      category=category, branch_label=city + ' · ' + address[:70], data_notes='؛ '.join(issues),
                      contract_relationships=[dict(id='contract_staff_' + code, route='direct_we', network_members=['employee'], source_ref='staff-20261018', source_row=row_number, financial_terms_status='unknown')])
        if kind == 'records':
            entity.update(provider_id='provider_staff_' + code, name='شريف موافي', display_name='د/ شريف موافي', raw_name=name, raw_phone=valid_phone, specialty='طب وجراحة الأسنان', hospital='تعاقد مباشر من الشركة WE', is_department=False, clinic_type='عيادة خارجية', category='doctors', schedule='المواعيد تحتاج تأكيدًا')
        for brand, (filename, url) in LOGOS.items():
            if brand in name and (ROOT / 'img/providers' / filename).exists():
                entity.update(image_kind='logo', image_url='https://abomawadda.github.io/Medical/img/providers/' + filename, image_source_url=url, image_verified=True)
        if kind == 'hospitals' and code in ALIASES:
            entity['short_name'] = {'758':'مستشفى السلاب', '766':'مستشفى الخير', '767':'مستشفى الحكمة'}[code]
            entity['phone_digits'] = next(iter(contact['landline'] + contact['mobile']), '')
            entity['hotline'] = ''  # Do not retain a conflicting historical number.
            entity['working_hours'] = 'المواعيد تحتاج تأكيدًا مع الجهة'
        if status == 'جاهز للدمج':
            accepted.append(dict(kind=kind, entity=entity))
        stats[status] += 1
        for issue in issues:
            reviews.append([code, row_number, name, status, issue, 'مفتوح'])
        normalized.append([code, name, category, governorate, city, markaz, address, '\n'.join(contact['mobile']), '\n'.join(contact['landline']), '\n'.join(contact['shortcode']), location, status, '؛ '.join(issues)])
        for k in ['mobile', 'landline', 'shortcode']:
            for n in contact[k]:
                checked = EVIDENCE.get(n)
                phones.append([code, name, {'mobile':'محمول','landline':'أرضي','shortcode':'مختصر'}[k], n, 'مطابق لمصدر منشور؛ لم يُختبر بمكالمة' if checked else 'صيغة سليمة؛ يحتاج تأكيد الجهة', checked[0] if checked else '', DATE if checked else ''])
        original.append([row_number] + [str(v or '') for v in row])
    summary = dict(source_rows=len(rows), imported_rows=len(accepted), quarantined_rows=stats['معلق للمراجعة'], outside_scope_rows=stats['خارج النطاق الحالي'], invalid_phone_values=stats['invalid_phones'], corrections=len(corrections), source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), date=DATE)
    IMPORT.parent.mkdir(exist_ok=True)
    IMPORT.write_text(json.dumps(dict(schema_version=1, summary=summary, entries=accepted), ensure_ascii=False, indent=2))
    output = Path(args.output_dir)
    output.mkdir(parents=True, exist_ok=True)
    result = openpyxl.Workbook()
    result.remove(result.active)
    sheet(result, 'ملخص', ['البند','القيمة'], [[k,v] for k,v in summary.items()] + [['ملاحظة','سليم الصيغة لا يعني أن الرقم يعمل؛ لم تنفذ مكالمات. الفوترة الإلكترونية وonline/offline ليست قواعد أهلية أو حجز.']])
    headers = ['كود الجهة','اسم مقدم الخدمة','النوع','المحافظة','المدينة أو القرية','المركز','العنوان','محمول','أرضي','مختصر','الخريطة','حالة الدمج','ملاحظات']
    sheet(result, 'جاهز للدمج', headers, [r for r in normalized if r[-2]=='جاهز للدمج'])
    sheet(result, 'معلق للمراجعة', headers, [r for r in normalized if r[-2]=='معلق للمراجعة'])
    sheet(result, 'خارج الدقهلية', headers, [r for r in normalized if r[-2]=='خارج النطاق الحالي'])
    sheet(result, 'أرقام الاتصال', ['كود الجهة','الاسم','الفئة','الرقم كنص','حالة التحقق','المصدر','تاريخ المطابقة'], phones)
    sheet(result, 'علاقات التعاقد', ['معرّف العلاقة','كود الجهة','اسم الجهة','المستفيد','مسار التعاقد','الشروط المالية','صف المصدر'], [[e['entity']['contract_relationships'][0]['id'],e['entity']['staff_provider_code'],e['entity']['name'],'الموظف','مباشر من الشركة حسب القائمة','لم تُحدد بعد',e['entity']['source_row']] for e in accepted])
    sheet(result, 'الملاحظات', ['كود الجهة','صف المصدر','الاسم','حالة الدمج','الملاحظة','حالة المراجعة'], reviews)
    sheet(result, 'سجل التصحيحات', ['كود الجهة','صف المصدر','الحقل','قبل','بعد','التفسير'], corrections)
    result.save(output / 'staff-network-organized.xlsx')
    # Full source is retained locally only, separately from the shareable workbook.
    private = openpyxl.Workbook()
    private.remove(private.active)
    sheet(private,'الأصل للمراجعة فقط',['صف المصدر']+[str(v or '') for v in list(book.active.values)[0]],original)
    private.save(output / 'staff-source-preserved-private.xlsx')
    (output/'import-summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
